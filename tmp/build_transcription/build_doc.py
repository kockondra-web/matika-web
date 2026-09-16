from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUT = Path(r"C:\Users\42073\Desktop\Matematika\web\output\documents\exponencialni_rovnice_6_9_2024.docx")
OUT.parent.mkdir(parents=True, exist_ok=True)


def set_font(run, name="Aptos", size=11, bold=False, color="111111"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=140, bottom=110, end=140):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def _parts(value):
    if isinstance(value, str):
        return [mr(value)]
    if hasattr(value, "tag"):
        return [value]
    result = []
    for item in value:
        result.extend(_parts(item))
    return result


def mr(text):
    node = OxmlElement("m:r")
    props = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Cambria Math")
    fonts.set(qn("w:hAnsi"), "Cambria Math")
    props.append(fonts)
    node.append(props)
    t = OxmlElement("m:t")
    t.text = text
    node.append(t)
    return node


def sup(base, exponent):
    node = OxmlElement("m:sSup")
    e = OxmlElement("m:e")
    for part in _parts(base): e.append(part)
    s = OxmlElement("m:sup")
    for part in _parts(exponent): s.append(part)
    node.extend((e, s))
    return node


def sub(base, subscript):
    node = OxmlElement("m:sSub")
    e = OxmlElement("m:e")
    for part in _parts(base): e.append(part)
    s = OxmlElement("m:sub")
    for part in _parts(subscript): s.append(part)
    node.extend((e, s))
    return node


def frac(numerator, denominator):
    node = OxmlElement("m:f")
    num = OxmlElement("m:num")
    den = OxmlElement("m:den")
    for part in _parts(numerator): num.append(part)
    for part in _parts(denominator): den.append(part)
    node.extend((num, den))
    return node


def radical(expression, degree=None):
    node = OxmlElement("m:rad")
    props = OxmlElement("m:radPr")
    hide = OxmlElement("m:degHide")
    hide.set(qn("m:val"), "1" if degree is None else "0")
    props.append(hide)
    deg = OxmlElement("m:deg")
    if degree is not None:
        for part in _parts(degree): deg.append(part)
    expr = OxmlElement("m:e")
    for part in _parts(expression): expr.append(part)
    node.extend((props, deg, expr))
    return node


def delim(expression, begin="(", end=")"):
    node = OxmlElement("m:d")
    props = OxmlElement("m:dPr")
    beg = OxmlElement("m:begChr")
    beg.set(qn("m:val"), begin)
    finish = OxmlElement("m:endChr")
    finish.set(qn("m:val"), end)
    props.extend((beg, finish))
    expr = OxmlElement("m:e")
    for part in _parts(expression): expr.append(part)
    node.extend((props, expr))
    return node


def log(base=None):
    return mr("log") if base is None else sub("log", base)


def eqp(container, expression, align=WD_ALIGN_PARAGRAPH.LEFT, left=0.28, after=2, size=12):
    p = container.add_paragraph()
    p.alignment = align
    p.paragraph_format.left_indent = Inches(left)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.keep_together = True
    math = OxmlElement("m:oMath")
    for part in _parts(expression): math.append(part)
    p._p.append(math)
    return p


E = {
    "a1": [sup("2", "2x-1"), " = ", radical(sup("8", "x+6"), "5")],
    "a2": [log("2"), sup("2", "2x-1"), " = ", log("2"), sup("2", frac("3x+18", "5"))],
    "a3": ["2x-1 = ", frac("3x+18", "5")],
    "b1": ["3·", sup("2", "x-1"), "+4·", sup("2", "x+1"), " = 38"],
    "b2": ["3·", sup("2", "x"), "·", sup("2", "-1"), "+4·", sup("2", "x"), "·", sup("2", "1"), " = 38"],
    "b3": [sup("2", "x"), delim([frac("3", "2"), "+8"]), " = 38"],
    "b4": [sup("2", "x"), " = 4"],
    "b5": ["x = 2"],
    "note2_1": [sup("2", "x"), " = 7"],
    "note2_2": [log("2"), sup("2", "x"), " = ", log("2"), "7"],
    "note2_3": ["x = ", log("2"), "7"],
    "noteln_1": ["ln ", sup("2", "x"), " = ln 7"],
    "noteln_2": ["x·ln 2 = ln 7"],
    "noteln_3": ["x = ", frac("ln 7", "ln 2")],
    "note10_1": ["log ", sup("2", "x"), " = log 7"],
    "note10_2": ["x = ", frac("log 7", "log 2")],
    "note10_3": ["x = ", frac("1", [log("7"), "2"])],
    "ex1": [sup("2", "x"), " = 5"],
    "ex2": [sup("4", "2x-1"), " = 10"],
    "ex3": [sup("4", "2-5x"), " = ", sup("5", "4x+1")],
    "ex4": [sup("2", "1-x"), "·", sup("3", "2x+3"), " = ", sup("5", "4-3x")],
    "ex5": [frac(radical(sup("4", "2x-5"), "3"), sup("9", "1-x")), " = ", sup("3", "5x+1"), "·", sup("8", "x-5")],
    "s1": ["log ", delim(frac(radical(sup("4", "2x-5"), "3"), sup("9", "1-x"))), " = log ", delim([sup("3", "5x+1"), "·", sup("8", "x-5")])],
    "s2": ["log ", radical(sup("4", "2x-5"), "3"), " - log ", sup("9", "1-x"), " = log ", sup("3", "5x+1"), " + log ", sup("8", "x-5")],
    "s3": [frac("2x-5", "3"), " log 4 - (1-x) log 9 = (5x+1) log 3 + (x-5) log 8"],
    "s4": [frac("4x-10", "3"), " log 2 - (2-2x) log 3 = (5x+1) log 3 + (3x-15) log 2     |·3"],
    "s5": ["4x log 2 - 10 log 2 - 6 log 3 + 6x log 3 = 15x log 3 + 3 log 3 + 9x log 2 - 45 log 2"],
    "s6": ["35 log 2 - 9 log 3 = 5x log 2 + 9x log 3"],
}

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.5), Inches(11)
section.top_margin = section.bottom_margin = Inches(0.68)
section.left_margin = section.right_margin = Inches(0.82)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
normal.font.size = Pt(11)
normal.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
normal.paragraph_format.space_after = Pt(5)
for style_name, size in (("Title", 22), ("Heading 1", 16), ("Heading 2", 13)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
styles["Title"].paragraph_format.space_after = Pt(9)
styles["Heading 1"].paragraph_format.space_before = Pt(10)
styles["Heading 1"].paragraph_format.space_after = Pt(5)
styles["Heading 1"].paragraph_format.keep_with_next = True

doc.add_paragraph("Exponenciální rovnice řešené logaritmováním", style="Title")
meta = doc.add_table(rows=2, cols=2)
meta.autofit = False
set_table_borders(meta, color="D9E2F3")
for row in meta.rows:
    row.cells[0].width, row.cells[1].width = Inches(1.6), Inches(5.1)
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell)
for idx, (label, value) in enumerate((("Hodiny", "1 a 2   |   6. 9. 2024"), ("Téma", "Seznámení s učivem, s pravidly klasifikace a s možnostmi konzultací"))):
    set_cell_shading(meta.cell(idx, 0), "D9E2F3")
    set_font(meta.cell(idx, 0).paragraphs[0].add_run(label), bold=True)
    set_font(meta.cell(idx, 1).paragraphs[0].add_run(value))
p = doc.add_paragraph()
p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(8), Pt(7)
set_font(p.add_run("Přepis zápisu z tabule zachovává uvedené příklady, pomocné poznámky a rozepsaný postup řešení."), size=10, color="4F4F4F")

doc.add_paragraph("Ukázka A", style="Heading 1")
eqp(doc, E["a1"], size=14)
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.28)
set_font(p.add_run("Zlogaritmujeme se základem 2:"), size=10, bold=True, color="C00000")
eqp(doc, E["a2"], size=14)
eqp(doc, E["a3"], size=14, after=5)

doc.add_paragraph("Ukázka B", style="Heading 1")
for name in ("b1", "b2", "b3", "b4", "b5"):
    eqp(doc, E[name], size=14)

doc.add_page_break()
doc.add_paragraph("Tři zápisy stejného řešení", style="Heading 1")
p = doc.add_paragraph()
set_font(p.add_run("Pro rovnici 2ˣ = 7 lze použít různé základy logaritmu."), size=10, color="4F4F4F")
methods = doc.add_table(rows=1, cols=3)
methods.autofit = False
methods_data = (
    ("Logaritmus se základem 2", "FCE4E4", "C00000", ("note2_1", "note2_2", "note2_3")),
    ("Přirozený logaritmus", "E2F0D9", "008000", ("noteln_1", "noteln_2", "noteln_3")),
    ("Dekadický logaritmus", "E7E6E6", "111111", ("note10_1", "note10_2", "note10_3")),
)
for i, (heading, fill, color, names) in enumerate(methods_data):
    cell = methods.cell(0, i)
    cell.width = Inches(2.1)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=130, start=90, bottom=130, end=90)
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(para.add_run(heading), size=9, bold=True, color=color)
    for name in names:
        eqp(cell, E[name], align=WD_ALIGN_PARAGRAPH.CENTER, left=0, after=1, size=11)
set_table_borders(methods)

doc.add_paragraph("Cvičení A", style="Heading 1")
for index, name in enumerate(("ex1", "ex2", "ex3", "ex4", "ex5"), start=1):
    if index == 5:
        doc.add_page_break()
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.cell(0, 0).width, table.cell(0, 1).width = Inches(0.35), Inches(6.15)
    for cell in table.rows[0].cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(cell, top=35, start=30, bottom=35, end=30)
    set_font(table.cell(0, 0).paragraphs[0].add_run(f"{index})"), bold=True)
    eqp(table.cell(0, 1), E[name], left=0, after=0, size=13)

doc.add_paragraph("Řešení úlohy 5", style="Heading 1")
p = doc.add_paragraph()
set_font(p.add_run("a) Rovnou logaritmujeme."), bold=True)
for name in ("s1", "s2", "s3", "s4", "s5", "s6"):
    eqp(doc, E[name], left=0.05, after=7, size=11 if name == "s5" else 12)

doc.core_properties.title = "Exponenciální rovnice řešené logaritmováním"
doc.core_properties.subject = "Přepis zápisu z hodin matematiky 6. 9. 2024"
doc.core_properties.author = ""
doc.save(OUT)
print(OUT)
